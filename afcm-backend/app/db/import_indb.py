"""
Builds a foods CSV from the Indian Nutrient Databank (INDB.xlsx), an open-access spreadsheet of
~1,014 Indian recipes with computed nutrition (github.com/lindsayjaacks/Indian-Nutrient-Databank-INDB-).

Unlike USDA, INDB gives dish-level values directly (e.g. "Chapati", "Chicken biryani"), which is
what Khana's photo classes need. Its columns are already per 100 g, in kcal, so no unit conversion
is needed, unlike the raw-ingredient IFCT2017 book.

1. Find the INDB recipe that matches each of your food ids:
       python -m app.db.import_indb --file INDB.xlsx --search "chicken biryani"

2. Write a mapping file (indb_map.csv). Only `id` and `recipe_code` are required:
       id,recipe_code,name,reference_weight_g,aliases
       chapati,ASC123,Chapati,40,roti|phulka

3. Turn it into a foods CSV, review it, then import it:
       python -m app.db.import_indb --file INDB.xlsx --map indb_map.csv --out foods_indb.csv
       python -m app.db.import_foods foods_indb.csv --dry-run
       python -m app.db.import_foods foods_indb.csv --update

INDB is described by its authors as open access; check its GitHub repository for the exact terms
before shipping a product built on it.
"""
import argparse
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

# INDB column -> our foods CSV column. Values are already per 100 g.
FIELD_MAP = {
    "energy_kcal": "calories",
    "protein_g": "protein_g",
    "carb_g": "carbs_g",
    "fat_g": "fat_g",
    "fibre_g": "fiber_g",
    "freesugar_g": "sugar_g",
    "sodium_mg": "sodium_mg",
    "cholesterol_mg": "cholesterol_mg",
}
# INDB reports saturated fat in milligrams as sfa_mg; our column is grams.
SATURATED_FAT_MG_COLUMN = "sfa_mg"
OUT_HEADER = [
    "id", "name", "calories", "reference_weight_g", "aliases", "protein_g", "carbs_g", "fat_g",
    "fiber_g", "sugar_g", "saturated_fat_g", "sodium_mg", "cholesterol_mg", "kind",
]


def load_recipes(path: Path) -> dict[str, dict]:
    """recipe_code -> {column name: value}, reading only the header + data (fast, read-only mode)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    if "food_code" not in header or "food_name" not in header:
        raise ValueError(f"{path.name} doesn't look like INDB.xlsx (expected food_code and food_name columns).")

    recipes: dict[str, dict] = {}
    for values in rows:
        record = dict(zip(header, values))
        code = record.get("food_code")
        if code:
            recipes[str(code).strip()] = record
    wb.close()
    return recipes


def search(path: Path, terms: str, limit: int = 15) -> list[tuple[str, str]]:
    words = [w for w in terms.lower().split() if w]
    hits = []
    for code, record in load_recipes(path).items():
        name = str(record.get("food_name") or "")
        if all(w in name.lower() for w in words):
            hits.append((len(name), code, name))
    hits.sort()
    return [(code, name) for _, code, name in hits[:limit]]


def _num(value: object) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def build_rows(path: Path, mapping: Path) -> tuple[list[dict], list[str]]:
    wanted = list(csv.DictReader(mapping.open(newline="", encoding="utf-8-sig")))
    for i, row in enumerate(wanted, start=2):
        if not (row.get("id") or "").strip() or not (row.get("recipe_code") or "").strip():
            raise ValueError(f"{mapping.name} line {i}: id and recipe_code are both required")

    recipes = load_recipes(path)
    rows: list[dict] = []
    warnings: list[str] = []
    for row in wanted:
        food_id, code = row["id"].strip().lower(), row["recipe_code"].strip()
        record = recipes.get(code)
        if record is None:
            raise ValueError(f"recipe_code {code} (for {food_id}) is not in {path.name}. Search with --search.")

        calories = _num(record.get("energy_kcal"))
        if calories <= 0:
            raise ValueError(f"INDB recipe {code} ({record.get('food_name')}) has no energy value, so it can't be used for {food_id}.")

        values = {column: _num(record.get(source)) for source, column in FIELD_MAP.items()}
        values["saturated_fat_g"] = _num(record.get(SATURATED_FAT_MG_COLUMN)) / 1000  # mg -> g

        macro_kcal = 4 * values["protein_g"] + 4 * values["carbs_g"] + 9 * values["fat_g"]
        if calories > 20 and not (0.5 * calories <= macro_kcal <= 1.5 * calories):
            warnings.append(f"{food_id}: calories ({calories:g}) and macros (about {macro_kcal:.0f} kcal) disagree. Check the recipe_code.")

        rows.append(
            {
                "id": food_id,
                "name": (row.get("name") or "").strip() or str(record.get("food_name") or food_id),
                "reference_weight_g": (row.get("reference_weight_g") or "").strip() or "100",
                "aliases": (row.get("aliases") or "").strip(),
                "kind": "food",
                **{c: f"{v:g}" for c, v in values.items()},
            }
        )
    return rows, warnings


def write_csv(rows: list[dict], out: Path) -> None:
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUT_HEADER)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--file", type=Path, required=True, help="path to INDB.xlsx")
    p.add_argument("--search", help="words to look for in the recipe name")
    p.add_argument("--map", type=Path, help="CSV with columns id, recipe_code (and optionally name, reference_weight_g, aliases)")
    p.add_argument("--out", type=Path, default=Path("foods_indb.csv"))
    args = p.parse_args()

    try:
        if args.search:
            hits = search(args.file, args.search)
            if not hits:
                print("No matches. Try fewer or different words.")
            for code, name in hits:
                print(f"{code:>10}  {name}")
            return 0
        if not args.map:
            p.error("give --search or --map")
        rows, warnings = build_rows(args.file, args.map)
    except (ValueError, OSError, KeyError) as exc:
        sys.exit(f"Error: {exc}")

    write_csv(rows, args.out)
    for w in warnings:
        print(f"WARNING  {w}")
    print(f"Wrote {len(rows)} foods to {args.out}. Review it, then: python -m app.db.import_foods {args.out} --dry-run")
    return 0


if __name__ == "__main__":
    sys.exit(main())